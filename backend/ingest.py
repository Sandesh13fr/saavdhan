"""Local ingestion: structured records are deterministic; scans fail visibly."""
from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path
from PIL import Image

from .models import Invoice, Payment, SourceRef, Supplier, Buyer, RatePeriod, Message, PurchaseOrder, Delivery, Allocation

MAX_FILE_BYTES = 15 * 1024 * 1024
ALLOWED = {".csv", ".xlsx", ".pdf", ".txt", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}


def tesseract_command():
    if os.getenv("BAKAYA_OCR_ENABLED", "false").lower() != "true":
        return None
    configured = os.getenv("TESSERACT_CMD")
    executable = shutil.which(configured or "tesseract")
    if not executable and Path("C:/Program Files/Tesseract-OCR/tesseract.exe").is_file():
        executable = None if configured else "C:/Program Files/Tesseract-OCR/tesseract.exe"
    return executable


def ocr_available():
    return bool(tesseract_command())


def ocr_image(data: bytes) -> str:
    """Local, opt-in OCR produces reviewable text, never financial facts."""
    executable = tesseract_command()
    if not executable:
        return ""
    with Image.open(io.BytesIO(data)) as image:
        if getattr(image, "n_frames", 1) != 1:
            raise ValueError("Multi-frame images require a page-preserving PDF upload")
        if image.width * image.height > 20_000_000:
            raise ValueError("OCR image exceeds 20 million pixels")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "page.png"
            image.convert("RGB").save(path)
            result = subprocess.run([executable, str(path), "stdout", "-l", "eng", "--psm", "6"],
                                    capture_output=True, timeout=30, check=True)
    return result.stdout.decode("utf-8")[:100000].strip()


def parse_date(value):
    # Deliberately do not guess ambiguous dd/mm vs mm/dd dates.
    return date.fromisoformat(str(value).strip()[:10])


def parse_money(value):
    return Decimal(str(value).replace(",", "").replace("₹", "").replace("INR", "").strip())


def row_record(row, sid):
    row = {str(k).strip().lower(): str(v).strip() for k, v in row.items() if k is not None and v is not None and str(v).strip()}
    kind = row.pop("record_type", row.pop("type", ""))
    kwargs = {"source_refs": [sid], "field_sources": {k:[sid] for k in row}}
    if kind == "invoice" or ("invoice_number" in row and "amount" in row):
        if "invoice_number" in row:
            row["number"] = row.pop("invoice_number")
            kwargs["field_sources"]["number"]=[sid]
        if "invoice_date" in row:
            row["date"] = row.pop("invoice_date")
            kwargs["field_sources"]["date"]=[sid]
        row.setdefault("id", row.get("number", f"INV-{uuid.uuid4().hex[:8]}"))
        for key in ("date","delivery_date","acceptance_date","contract_date","objection_date","objection_resolved_date"):
            if key in row:
                row[key]=parse_date(row[key])
        row["amount"]=parse_money(row["amount"])
        if "written_term_days" in row:
            row["written_term_days"]=int(row["written_term_days"])
        return "invoices",Invoice(**kwargs,**row)
    if kind == "payment" or ("date" in row and "amount" in row):
        row.setdefault("id", f"PAY-{uuid.uuid4().hex[:8]}")
        row["date"]=parse_date(row["date"])
        row["amount"]=parse_money(row["amount"])
        for key in ("invoice_ids",):
            if key in row:
                row[key]=[s.strip() for s in row[key].split("|") if s.strip()]
        if "tds_amount" in row:
            row["tds_amount"]=parse_money(row["tds_amount"])
        return "payments",Payment(**kwargs,**row)
    if kind == "supplier":
        if "udyam_date" in row:
            row["udyam_date"]=parse_date(row["udyam_date"])
        return "supplier",Supplier(**kwargs,**row)
    if kind == "buyer":
        return "buyer",Buyer(**kwargs,**row)
    if kind == "rate":
        for key in ("start","end"):
            row[key]=parse_date(row[key])
        row["bank_rate"]=Decimal(row["bank_rate"])
        if "synthetic" in row:
            row["synthetic"]=row["synthetic"].lower() in {"true","1","yes"}
        return "rates",RatePeriod(**kwargs,**row)
    if kind == "delivery":
        row["date"]=parse_date(row["date"])
        if "accepted" in row:
            row["accepted"]=row["accepted"].lower() in {"true","1","yes"}
        return "deliveries",Delivery(**kwargs,**row)
    if kind == "purchase_order":
        row["date"]=parse_date(row["date"])
        return "purchase_orders",PurchaseOrder(**kwargs,**row)
    if kind == "bank_control":
        return "bank_total",parse_money(row["amount"])
    raise ValueError("Unrecognized row. Include record_type and ISO dates; see sample CSV templates.")


def ingest_bytes(filename: str, data: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED:
        raise ValueError(f"Unsupported extension {suffix}; accepted: CSV, XLSX, PDF, TXT, images")
    if len(data)>MAX_FILE_BYTES:
        raise ValueError("File exceeds15 MB limit")
    root = "SRC-"+uuid.uuid4().hex[:12]
    base = SourceRef(id=root,file=Path(filename).name,kind="uploaded_document",sha256=hashlib.sha256(data).hexdigest(),stored_name=root+suffix)
    sources, extracted, errors = [base], {}, []
    if suffix in {".png",".jpg",".jpeg",".webp",".tif",".tiff"}:
        base.status="ocr_unavailable"
        base.text="OCR is not configured. Re-upload text-searchable evidence or provide a structured CSV with verified facts. No date or amount inferred."
        try:
            text = ocr_image(data)
            if text:
                base.text, base.status = text, "unstructured_review"
                base.kind = "ocr_document"
        except (OSError, ValueError, subprocess.SubprocessError, Image.DecompressionBombError):
            errors.append("OCR failed or exceeded its limit; upload searchable evidence or verified structured facts.")
        return sources,extracted,errors
    if suffix in {".csv",".xlsx"}:
        if suffix == ".csv":
            text=data.decode("utf-8-sig")
            table=list(csv.DictReader(io.StringIO(text)))
        else:
            from openpyxl import load_workbook
            workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            sheet=workbook.active
            rows=list(sheet.iter_rows(values_only=True))
            if not rows:
                raise ValueError("Workbook is empty")
            headers=[str(v or "").strip() for v in rows[0]]
            table=[dict(zip(headers,row)) for row in rows[1:]]
            workbook.close()
            text="\n".join(",".join(str(v or "") for v in row.values()) for row in table)
        base.text=text[:100000]
        base.kind="structured_ledger"
        for n,row in enumerate(table,2):
            sid=f"{root}-L{n}"
            source=base.model_copy(update={"id":sid,"line":n,"text":str(row),"stored_name":base.stored_name})
            sources.append(source)
            try:
                key,record=row_record(row,sid)
                if key in {"supplier","buyer","bank_total"}:
                    extracted[key]=record
                else:
                    extracted.setdefault(key,[]).append(record)
            except Exception as exc:
                source.status="parse_error"
                errors.append(f"Row{n}: {exc}")
        return sources,extracted,errors
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if len(reader.pages) > 50:
            raise ValueError("PDF exceeds 50 page limit")
        pages=[]
        for n,page in enumerate(reader.pages,1):
            text=page.extract_text() or ""
            if not text.strip() and os.getenv("BAKAYA_OCR_ENABLED", "false").lower() == "true":
                try:
                    import pypdfium2 as pdfium
                    with pdfium.PdfDocument(data) as document:
                        scan = document[n-1]
                        try:
                            if scan.get_width() * scan.get_height() * 9 > 20_000_000:
                                raise ValueError("OCR page exceeds pixel limit")
                            bitmap = scan.render(scale=3)
                            image = bitmap.to_pil()
                            output = io.BytesIO()
                            image.save(output, format="PNG")
                            text = ocr_image(output.getvalue())
                            image.close()
                            bitmap.close()
                        finally:
                            scan.close()
                except (OSError, ValueError, subprocess.SubprocessError, pdfium.PdfiumError):
                    errors.append(f"Page {n}: OCR failed; manual evidence review required.")
            sid=f"{root}-P{n}"
            source=base.model_copy(update={"id":sid,"page":n,"text":text[:100000],"stored_name":base.stored_name})
            source.status="unstructured_review" if text.strip() else "ocr_unavailable"
            if text.strip() and not (page.extract_text() or "").strip():
                source.kind = "ocr_document"
            sources.append(source)
            pages.append(text)
        base.text="\n".join(pages)[:100000]
        base.status="parsed" if base.text.strip() else "ocr_unavailable"
        return sources,extracted,errors
    text=data.decode("utf-8-sig")
    base.text=text[:100000]
    base.kind="message_export"
    # TXT is evidence only: message text cannot prove delivery/acceptance automatically.
    extracted["messages"]=[]
    for n,line in enumerate(text.splitlines(),1):
        if not line.strip():
            continue
        sid=f"{root}-L{n}"
        sources.append(base.model_copy(update={"id":sid,"line":n,"text":line}))
        extracted["messages"].append(Message(id=f"MSG-{root}-{n}",text=line,source_refs=[sid],field_sources={"text":[sid]}))
    return sources,extracted,errors
